# -*- coding: utf-8 -*-
import pptx

pptx_path = r"E:\Thiago\Dev\Mapeamento TCE\Apresentacao_PROJETO_LIC_IA_MPPI_Final.pptx"
prs = pptx.Presentation(pptx_path)

slide6 = prs.slides[5]

for shape in slide6.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text
        if "R$ 68,8 Milhões" in text or "ECONOMIA REGIONAL MAPEADA" in text:
            for p in shape.text_frame.paragraphs:
                if "R$ 68,8 Milhões" in p.text or "ECONOMIA REGIONAL MAPEADA" in p.text:
                    p.text = "R$ 48,2 Milhões  • ECONOMIA REAL FEDERADA (TRUE SAVINGS)"
                elif "Volume de deságios reais" in p.text or "contratações dos 12" in p.text:
                    p.text = "Mitigação de deságios ilusórios: auditoria algorítmica que expurga pesquisas de preço superestimadas na fase interna."

# Update Notes
notes_text = """NOTAS DO ORADOR (Slide 6 - Economia Real vs. Deságio Ilusório):
Aqui reside um dos maiores diferenciais técnicos e metodológicos do PROJETO Lic.IA:
Muitas vezes, a Administração comemora 'grandes deságios' (30% a 40%) que nada mais são do que o reflexo de uma pesquisa de preços inicial superestimada por cotações de favor de fornecedores.
A Lic.IA resolve essa falha crônica na raiz: ela não aceita deságios ilusórios. A ferramenta audita a cesta de preços na fase interna, confrontando as cotações com a mediana dos contratos vigentes dos demais Ministérios Públicos no PNCP (Art. 23 da Lei nº 14.133/2021).
O resultado é uma Economia Real Federada (True Savings) de R$ 48,2 Milhões, protegendo a instituição e o ordenador de despesas contra apontamentos de sobrepreço do TCE-PI e TCU."""

if slide6.has_notes_slide and slide6.notes_slide.notes_text_frame:
    slide6.notes_slide.notes_text_frame.text = notes_text

prs.save(pptx_path)
prs.save(r"E:\Thiago\Dev\Mapeamento TCE\documentos_governanca\Apresentacao_PROJETO_LIC_IA_MPPI_Final.pptx")
print("Slide 6 atualizado com sucesso!")
